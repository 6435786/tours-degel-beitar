import openpyxl
from openpyxl.utils import get_column_letter
import html, os, subprocess, re, sys
from pypdf import PdfReader

SRC = r"C:\Users\1\Desktop\תקציב ים סוף 14-16.xlsx"
OUT_HTML = r"C:\Users\1\Desktop\תקציב ים סוף 14-16.html"
OUT_PDF  = r"C:\Users\1\Desktop\תקציב ים סוף 14-16.pdf"
CHROME   = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

wb_vals = openpyxl.load_workbook(SRC, data_only=True)
wb_form = openpyxl.load_workbook(SRC, data_only=False)
SHEET_NAMES = wb_vals.sheetnames

def find_sheet_refs(formula):
    refs, seen = [], set()
    for m in re.finditer(r"'([^']+)'!", formula):
        nm = m.group(1)
        if nm not in seen and nm in SHEET_NAMES:
            refs.append(nm); seen.add(nm)
    for nm in SHEET_NAMES:
        if nm in seen:
            continue
        pat = re.compile(r"(?<!')" + re.escape(nm) + r"(?!')!")
        if pat.search(formula):
            refs.append(nm); seen.add(nm)
    return refs

def fmt_value(v):
    if v is None: return ""
    if isinstance(v, float):
        if v.is_integer(): return str(int(v))
        return f"{v:.4f}".rstrip('0').rstrip('.')
    return str(v)

def has_content(vcell, fcell):
    v, f = vcell.value, fcell.value
    if v is None and f is None: return False
    if isinstance(v, str) and v.strip() == "" and (f is None or (isinstance(f, str) and f.strip() == "")):
        return False
    return True

def build_html(sheet_to_page):
    parts = []
    parts.append("<!doctype html><html lang='he' dir='rtl'><head><meta charset='utf-8'>")
    parts.append("<title>תקציב ים סוף 14-16</title>")
    parts.append("""<style>
    @page { size: A3 landscape; margin: 22mm 8mm 18mm 8mm; }
    body { font-family: Arial, 'David', sans-serif; direction: rtl; font-size: 9px; margin: 0; }
    .sheet { page-break-before: always; }
    .sheet.first { page-break-before: auto; }
    table { border-collapse: collapse; table-layout: fixed; width: 100%; }
    thead { display: table-header-group; }
    th, td { border: 1px solid #9ab; padding: 2px 3px; vertical-align: top; text-align: right;
             word-wrap: break-word; overflow-wrap: break-word; word-break: break-word; }
    th { background: #d9e6f5; font-weight: bold; text-align: center; }
    .sheet-title { background: #2a5a9a; color: #fff !important; font-size: 15px; padding: 6px;
                   text-align: center; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
    .corner { background: #b8c7d9; font-weight: bold; text-align: center; }
    .rowhdr { background: #d9e6f5; font-weight: bold; text-align: center; }
    .val { display: block; }
    .formula { display: block; font-family: Consolas, monospace; font-size: 7.5px; color: #b00;
               background: #fff8e1; border-top: 1px dashed #c80; margin-top: 2px; padding-top: 1px;
               direction: ltr; text-align: left; unicode-bidi: embed; word-break: break-all; }
    .ref-annot { color: #0066a0; direction: rtl; display: block; font-size: 8px; font-weight: bold; }
    .num { text-align: left; direction: ltr; }
    .marker { color: white; font-size: 3pt; }
    </style></head><body>""")

    for idx, sheet_name in enumerate(SHEET_NAMES):
        ws_v = wb_vals[sheet_name]
        ws_f = wb_form[sheet_name]
        max_row = ws_v.max_row
        max_col = ws_v.max_column

        used_cols = []
        for c in range(1, max_col+1):
            for r in range(1, max_row+1):
                if has_content(ws_v.cell(r, c), ws_f.cell(r, c)):
                    used_cols.append(c); break
        used_rows = []
        for r in range(1, max_row+1):
            for c in used_cols:
                if has_content(ws_v.cell(r, c), ws_f.cell(r, c)):
                    used_rows.append(r); break
        if not used_cols or not used_rows:
            continue

        cls_first = " first" if idx == 0 else ""
        parts.append(f"<div class='sheet{cls_first}'>")
        parts.append(f"<span class='marker'>##PAGEMARK{idx}##</span>")

        n = len(used_cols)
        hdr_w = 2.5
        col_w = (100 - hdr_w) / n
        parts.append("<table><colgroup>")
        parts.append(f"<col style='width:{hdr_w}%'>")
        for _ in used_cols:
            parts.append(f"<col style='width:{col_w:.3f}%'>")
        parts.append("</colgroup>")

        parts.append("<thead>")
        parts.append(f"<tr><th class='sheet-title' colspan='{n+1}'>גליון: {html.escape(sheet_name)}</th></tr>")
        parts.append("<tr><th class='corner'>#</th>")
        for c in used_cols:
            parts.append(f"<th>{get_column_letter(c)}</th>")
        parts.append("</tr></thead><tbody>")

        for r in used_rows:
            parts.append(f"<tr><td class='rowhdr'>{r}</td>")
            for c in used_cols:
                vcell = ws_v.cell(r, c)
                fcell = ws_f.cell(r, c)
                val = vcell.value
                fv = fcell.value
                formula = None
                if isinstance(fv, str) and fv.startswith('='):
                    formula = fv
                elif fcell.data_type == 'f':
                    formula = str(fv) if fv is not None else None

                val_str = fmt_value(val)
                is_num = isinstance(val, (int, float))
                cls = "num" if is_num else ""
                cell_html = f"<td class='{cls}'>"
                if val_str != "":
                    cell_html += f"<span class='val'>{html.escape(val_str)}</span>"
                else:
                    cell_html += "<span class='val'>&nbsp;</span>"
                if formula:
                    cell_html += f"<span class='formula'>{html.escape(formula)}</span>"
                    for rname in find_sheet_refs(formula):
                        pnum = sheet_to_page.get(rname)
                        if pnum:
                            safe = html.escape(rname)
                            cell_html += f"<span class='ref-annot'>(הפניה לגליון '{safe}' – עמוד {pnum})</span>"
                cell_html += "</td>"
                parts.append(cell_html)
            parts.append("</tr>")
        parts.append("</tbody></table></div>")

    parts.append("</body></html>")
    return "".join(parts)

def render_pdf():
    header_template = (
        "<div style='font-size:14px;width:100%;text-align:center;"
        "padding:2px 8px;font-weight:bold;color:#000;font-family:Arial;'>"
        "<span>עמוד </span><span class='pageNumber'></span>"
        "<span> מתוך </span><span class='totalPages'></span>"
        "</div>"
    )
    footer_template = header_template
    file_url = "file:///" + OUT_HTML.replace("\\", "/")
    cmd = [
        CHROME, "--headless=new", "--disable-gpu",
        f"--print-to-pdf={OUT_PDF}",
        f"--header-template={header_template}",
        f"--footer-template={footer_template}",
        file_url,
    ]
    return subprocess.run(cmd, capture_output=True, text=True, errors='replace')

def find_sheet_pages():
    reader = PdfReader(OUT_PDF)
    mapping = {}
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        for m in re.finditer(r'PAGEMARK(\d+)', text):
            k = int(m.group(1))
            if k < len(SHEET_NAMES) and SHEET_NAMES[k] not in mapping:
                mapping[SHEET_NAMES[k]] = i + 1
    return mapping

# Pass 1 -- no annotations, just to locate each sheet's page number
with open(OUT_HTML, "w", encoding="utf-8") as f:
    f.write(build_html({}))
print("Pass 1: rendering to locate sheet pages...")
r = render_pdf()
if r.returncode != 0:
    print("Chrome failed:", r.stderr[-500:]); sys.exit(1)

sheet_to_page = find_sheet_pages()
print("Sheet page mapping:")
for k, v in sheet_to_page.items():
    print(f"  {k!r} -> page {v}")

# Pass 2 -- with annotations
with open(OUT_HTML, "w", encoding="utf-8") as f:
    f.write(build_html(sheet_to_page))
print("Pass 2: rendering final PDF with sheet references...")
r = render_pdf()
print("Done. PDF:", OUT_PDF, "exists:", os.path.exists(OUT_PDF))
